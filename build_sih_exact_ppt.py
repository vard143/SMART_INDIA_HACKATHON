"""
Builds the exact Smart India Hackathon 2026 Presentation matching the provided template layout,
with embedded graphic images, exact typography, team name oval, and featuring:
- JCERT Government Curriculum ('भाषा अंजलि', 'गणित ज्ञान')
- NIPUN Bharat Assessment & Examination Engine with Printable PDF Report Cards
"""

import os
import time
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path="SIH2026_Idea_Submission_PS26042.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette matching the template
    NAVY_TITLE = RGBColor(30, 58, 138)     # Deep Template Blue #1E3A8A
    HEADING_BLACK = RGBColor(15, 23, 42)    # Slate 900
    TEXT_BLACK = RGBColor(30, 41, 59)       # Body text
    MUTED_GRAY = RGBColor(100, 116, 139)
    BANNER_BLUE = RGBColor(2, 132, 199)    # Footer line color #0284C7
    PALASH_ORANGE = RGBColor(235, 56, 7)
    WHITE = RGBColor(255, 255, 255)

    def add_common_decorations(slide, slide_title, slide_num):
        # 1. Top-Left Oval: 'Your Team Name'
        oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.3), Inches(1.8), Inches(1.1))
        oval.fill.solid()
        oval.fill.fore_color.rgb = WHITE
        oval.line.color.rgb = HEADING_BLACK
        oval.line.width = Pt(1.5)
        
        otf = oval.text_frame
        otf.word_wrap = True
        op = otf.paragraphs[0]
        op.alignment = PP_ALIGN.CENTER
        op.text = "Your\nTeam\nName"
        op.font.name = 'Arial'
        op.font.size = Pt(11)
        op.font.bold = True
        op.font.color.rgb = HEADING_BLACK

        # 2. Top-Right: SIH 2026 Logo
        if os.path.exists("ppt_assets/sih_logo.png"):
            slide.shapes.add_picture("ppt_assets/sih_logo.png", Inches(10.6), Inches(0.2), width=Inches(2.3))

        # 3. Center Slide Heading
        tb_head = slide.shapes.add_textbox(Inches(2.4), Inches(0.4), Inches(8.0), Inches(0.8))
        htf = tb_head.text_frame
        hp = htf.paragraphs[0]
        hp.alignment = PP_ALIGN.CENTER
        hp.text = slide_title
        hp.font.name = 'Times New Roman'
        hp.font.size = Pt(28)
        hp.font.bold = True
        hp.font.color.rgb = HEADING_BLACK

        # 4. Bottom Footer Banner
        footer_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.0), Inches(13.333), Inches(0.5))
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = BANNER_BLUE
        footer_line.line.color.rgb = BANNER_BLUE

        tb_foot = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12.333), Inches(0.5))
        ftf = tb_foot.text_frame
        fp = ftf.paragraphs[0]
        fp.text = "@SIH Idea submission- Template"
        fp.font.name = 'Arial'
        fp.font.size = Pt(11)
        fp.font.bold = True
        fp.font.color.rgb = WHITE

        # Slide number on bottom right
        tb_num = slide.shapes.add_textbox(Inches(12.0), Inches(7.0), Inches(1.0), Inches(0.5))
        ntf = tb_num.text_frame
        np = ntf.paragraphs[0]
        np.alignment = PP_ALIGN.RIGHT
        np.text = str(slide_num)
        np.font.name = 'Arial'
        np.font.size = Pt(12)
        np.font.bold = True
        np.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 1: TITLE PAGE
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)

    # Top Header: SMART INDIA HACKATHON 2026
    tb1_top = slide1.shapes.add_textbox(Inches(1.0), Inches(0.35), Inches(9.5), Inches(0.8))
    tf1_top = tb1_top.text_frame
    p1_top = tf1_top.paragraphs[0]
    p1_top.text = "SMART INDIA HACKATHON 2026"
    p1_top.font.name = 'Times New Roman'
    p1_top.font.size = Pt(30)
    p1_top.font.bold = True
    p1_top.font.color.rgb = NAVY_TITLE

    # Top Right SIH Logo
    if os.path.exists("ppt_assets/sih_logo.png"):
        slide1.shapes.add_picture("ppt_assets/sih_logo.png", Inches(10.6), Inches(0.2), width=Inches(2.3))

    # Center Subheader: TITLE PAGE
    tb1_sub = slide1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.333), Inches(0.6))
    tf1_sub = tb1_sub.text_frame
    p1_sub = tf1_sub.paragraphs[0]
    p1_sub.alignment = PP_ALIGN.CENTER
    p1_sub.text = "TITLE PAGE"
    p1_sub.font.name = 'Times New Roman'
    p1_sub.font.size = Pt(24)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = HEADING_BLACK

    # Left: Content Bullets matching the exact template format
    tb1_left = slide1.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(7.5), Inches(4.8))
    tf1_left = tb1_left.text_frame
    tf1_left.word_wrap = True

    bullets1 = [
        ("• Problem Statement ID – ", "26042"),
        ("• Problem Statement Title - ", "AI-Powered Vernacular Pedagogy and Real-Time Translation Tool for Mother Tongue-Based Primary Education (BHASHASETU)"),
        ("• Theme - ", "Smart Education"),
        ("• PS Category - ", "Software"),
        ("• Team ID - ", "[Your Team ID]"),
        ("• Team Name (Registered on portal) - ", "[Your Registered Team Name]")
    ]

    for idx, (label, val) in enumerate(bullets1):
        p = tf1_left.paragraphs[0] if idx == 0 else tf1_left.add_paragraph()
        p.space_after = Pt(14)
        r1 = p.add_run()
        r1.text = label
        r1.font.name = 'Arial'
        r1.font.bold = True
        r1.font.size = Pt(15)
        r1.font.color.rgb = HEADING_BLACK

        r2 = p.add_run()
        r2.text = val
        r2.font.name = 'Arial'
        r2.font.bold = (idx <= 3)
        r2.font.size = Pt(15)
        r2.font.color.rgb = NAVY_TITLE if idx <= 3 else MUTED_GRAY

    # Right: Lightbulb Brain Graphic from Template
    if os.path.exists("ppt_assets/slide1_bulb.png"):
        slide1.shapes.add_picture("ppt_assets/slide1_bulb.png", Inches(8.4), Inches(2.0), width=Inches(4.5))

    # ==========================================
    # SLIDE 2: IDEA TITLE & PROPOSED SOLUTION
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_common_decorations(slide2, "IDEA TITLE: BHASHASETU (भाषा सेतु (BhashaSetu))", 2)

    # Sub-heading with diamond symbol
    tb2_sub = slide2.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.0), Inches(0.5))
    tf2_sub = tb2_sub.text_frame
    p2_sub = tf2_sub.paragraphs[0]
    p2_sub.text = "❖ Proposed Solution (Describe your Idea/Solution/Prototype)"
    p2_sub.font.name = 'Arial'
    p2_sub.font.size = Pt(18)
    p2_sub.font.bold = True
    p2_sub.font.color.rgb = NAVY_TITLE

    # Left: Core Content Bullets
    tb2_left = slide2.shapes.add_textbox(Inches(0.6), Inches(1.85), Inches(6.8), Inches(5.0))
    tf2_left = tb2_left.text_frame
    tf2_left.word_wrap = True

    sections2 = [
        ("• Detailed explanation of the proposed solution:", [
            "Live Voice Bridge (<1.2s): Real-time two-way voice translation between Hindi teachers & tribal students (Ho, Mundari, Santhali) with sub-1.2s latency (<3s mandate).",
            "JCERT Government Curriculum Hub: Aligned with Jharkhand State textbooks ('भाषा अंजलि', 'गणित ज्ञान', 'हमारा परिवेश') for Balvatika to Class 3 with Ol Chiki (ᱚᱞ ᱪᱤᱠᱤ) & Devanagari scripts.",
            "NIPUN Assessment & Exam Hub: Periodic formative & summative FLN exams with instant AI auto-grading, oral reading fluency scoring, and 1-click printable PDF report cards.",
            "Bilingual Worksheet Studio: 1-click auto-generated printable vector A4 PDF activity sheets (matching, tracing, counting with tribal realia)."
        ]),
        ("• How it addresses the problem:", [
            "Bridges the severe language gap for 20,000+ Hindi-trained teachers across 5,000+ primary schools without requiring prior tribal language training.",
            "Enforces NIPUN Bharat FLN learning outcomes and standardizes continuous comprehensive evaluation in vernacular medium."
        ]),
        ("• Innovation and uniqueness of the solution:", [
            "100% Offline Edge Operation: Functions seamlessly on low-cost (>=2GB RAM, Android 9+) tablets with ~135MB RAM footprint and zero internet required.",
            "Integrated Government Curriculum + Diagnostic Exam Engine: Combines real-time translation with official JCERT syllabus and competency tracking."
        ])
    ]

    is_first = True
    for sec_title, bullet_items in sections2:
        p = tf2_left.paragraphs[0] if is_first else tf2_left.add_paragraph()
        is_first = False
        p.space_after = Pt(3)
        r = p.add_run()
        r.text = sec_title
        r.font.name = 'Arial'
        r.font.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = HEADING_BLACK

        for item in bullet_items:
            pi = tf2_left.add_paragraph()
            pi.space_after = Pt(2)
            r_dot = pi.add_run()
            r_dot.text = "   - "
            r_dot.font.bold = True
            r_dot.font.size = Pt(10)
            r_dot.font.color.rgb = PALASH_ORANGE

            r_txt = pi.add_run()
            r_txt.text = item
            r_txt.font.size = Pt(10)
            r_txt.font.color.rgb = TEXT_BLACK

    # Right: Embedded Infographic Image
    if os.path.exists("ppt_assets/slide2_solution.png"):
        slide2.shapes.add_picture("ppt_assets/slide2_solution.png", Inches(7.6), Inches(1.85), width=Inches(5.2))

    # ==========================================
    # SLIDE 3: TECHNICAL APPROACH
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_common_decorations(slide3, "TECHNICAL APPROACH", 3)

    # Left: Content Bullets
    tb3_left = slide3.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(6.8), Inches(5.3))
    tf3_left = tb3_left.text_frame
    tf3_left.word_wrap = True

    sections3 = [
        ("• Technologies to be used:", [
            "Frontend / PWA: React 19, TypeScript, Vite, Tailwind CSS, Lucide Icons, Canvas Confetti (Optimized for 2GB Android Tablets).",
            "NLP & Speech Engine: Rule-Dictionary Morphological Engine, Web Speech API (STT), Web Audio Acoustic Phoneme Synthesizer (TTS).",
            "Assessment & PDF Generator: jsPDF + html2canvas (Printable high-contrast bilingual A4 worksheets & NIPUN Report Cards).",
            "Backend & Edge Sync: FastAPI (Python 3.14), Uvicorn, Pydantic, IndexedDB Local Cache & Service Worker Offline Storage."
        ]),
        ("• Methodology and process for implementation:", [
            "1. Speech Capture: Real-time dual mic input (Teacher Hindi / Student Tribal) with noise filtering.",
            "2. Fast Hybrid NLP (<5ms): Bidirectional dictionary lookup with fuzzy morphological tokenization and Ol Chiki (ᱚᱞ ᱪᱤᱠᱤ) script transliteration.",
            "3. Acoustic Synthesis (<1.2s): Procedural audio synthesizer generates instant tribal voice + teleprompter subtitles.",
            "4. Assessment & Exam Engine: Automatic scoring of oral reading fluency, listening comprehension, and numeracy with competency report generation.",
            "5. Edge Offline Storage: 100% offline functionality powered by IndexedDB with zero server dependency."
        ])
    ]

    is_first = True
    for sec_title, bullet_items in sections3:
        p = tf3_left.paragraphs[0] if is_first else tf3_left.add_paragraph()
        is_first = False
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = sec_title
        r.font.name = 'Arial'
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = HEADING_BLACK

        for item in bullet_items:
            pi = tf3_left.add_paragraph()
            pi.space_after = Pt(3)
            r_dot = pi.add_run()
            r_dot.text = "   - "
            r_dot.font.bold = True
            r_dot.font.size = Pt(10.5)
            r_dot.font.color.rgb = BANNER_BLUE

            r_txt = pi.add_run()
            r_txt.text = item
            r_txt.font.size = Pt(10.5)
            r_txt.font.color.rgb = TEXT_BLACK

    # Right: Embedded Flowchart Image
    if os.path.exists("ppt_assets/slide3_tech_flow.png"):
        slide3.shapes.add_picture("ppt_assets/slide3_tech_flow.png", Inches(7.6), Inches(1.5), width=Inches(5.2))

    # ==========================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_common_decorations(slide4, "FEASIBILITY AND VIABILITY", 4)

    # Left: Content Bullets
    tb4_left = slide4.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(6.8), Inches(5.3))
    tf4_left = tb4_left.text_frame
    tf4_left.word_wrap = True

    sections4 = [
        ("• Analysis of the feasibility of the idea:", [
            "Technical Feasibility: Fully validated working prototype achieving sub-1.2s latency, ~135MB RAM footprint, and instant exam evaluation on Android 9+ tablets.",
            "Operational Feasibility: Zero teacher training needed; 1-tap push-to-talk Walkie-Talkie and pre-built JCERT lesson plans.",
            "Economic Feasibility: 100% open-source stack with zero recurring cloud API costs."
        ]),
        ("• Potential challenges and risks:", [
            "Limited digital NLP datasets for low-resource tribal languages (Ho, Mundari, Santhali).",
            "Zero internet connectivity in remote forest tribal schools.",
            "Standardizing assessments across diverse tribal dialects and scripts."
        ]),
        ("• Strategies for overcoming these challenges:", [
            "Built hybrid rule-morphological parser with 2,000+ curated FLN terms & authentic script mapping.",
            "100% offline client-side IndexedDB caching ensures zero dependence on network connectivity.",
            "Bilingual multi-modal assessment engine supporting both Ol Chiki and Devanagari test formats."
        ])
    ]

    is_first = True
    for sec_title, bullet_items in sections4:
        p = tf4_left.paragraphs[0] if is_first else tf4_left.add_paragraph()
        is_first = False
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = sec_title
        r.font.name = 'Arial'
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = HEADING_BLACK

        for item in bullet_items:
            pi = tf4_left.add_paragraph()
            pi.space_after = Pt(3)
            r_dot = pi.add_run()
            r_dot.text = "   - "
            r_dot.font.bold = True
            r_dot.font.size = Pt(10.5)
            r_dot.font.color.rgb = PALASH_ORANGE

            r_txt = pi.add_run()
            r_txt.text = item
            r_txt.font.size = Pt(10.5)
            r_txt.font.color.rgb = TEXT_BLACK

    # Right: Embedded Feasibility Matrix Image
    if os.path.exists("ppt_assets/slide4_matrix.png"):
        slide4.shapes.add_picture("ppt_assets/slide4_matrix.png", Inches(7.6), Inches(1.5), width=Inches(5.2))

    # ==========================================
    # SLIDE 5: IMPACT AND BENEFITS
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_common_decorations(slide5, "IMPACT AND BENEFITS", 5)

    # Left: Content Bullets
    tb5_left = slide5.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(6.8), Inches(5.3))
    tf5_left = tb5_left.text_frame
    tf5_left.word_wrap = True

    sections5 = [
        ("• Potential impact on the target audience:", [
            "5,000+ Primary Schools: Deployed across Jharkhand's tribal heartland (Kolhan, Santhal Parganas, Ranchi, Khunti).",
            "100,000+ Tribal Learners: Continuous assessment and mother-tongue instruction eliminate exam fear and learning deficits.",
            "20,000+ Teachers: Automated NIPUN evaluation & lesson delivery save thousands of instructional hours."
        ]),
        ("• Benefits of the solution:", [
            "Social Benefits: Revitalizes endangered indigenous languages (Ho, Mundari, Santhali) and restores Ol Chiki (ᱚᱞ ᱪᱤᱠᱤ) script pride.",
            "Educational Benefits: Enforces NEP 2020 §4.11-4.14 mother-tongue mandate and measures actual NIPUN FLN competency gains with printable report cards.",
            "Economic Benefits: Runs on existing government tablets with zero cloud subscription or hardware replacement costs."
        ])
    ]

    is_first = True
    for sec_title, bullet_items in sections5:
        p = tf5_left.paragraphs[0] if is_first else tf5_left.add_paragraph()
        is_first = False
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = sec_title
        r.font.name = 'Arial'
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = HEADING_BLACK

        for item in bullet_items:
            pi = tf5_left.add_paragraph()
            pi.space_after = Pt(3)
            r_dot = pi.add_run()
            r_dot.text = "   - "
            r_dot.font.bold = True
            r_dot.font.size = Pt(10.5)
            r_dot.font.color.rgb = BANNER_BLUE

            r_txt = pi.add_run()
            r_txt.text = item
            r_txt.font.size = Pt(10.5)
            r_txt.font.color.rgb = TEXT_BLACK

    # Right: Embedded Impact Infographic Image
    if os.path.exists("ppt_assets/slide5_impact.png"):
        slide5.shapes.add_picture("ppt_assets/slide5_impact.png", Inches(7.6), Inches(1.5), width=Inches(5.2))

    # ==========================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_common_decorations(slide6, "RESEARCH AND REFERENCES", 6)

    # Full width cards
    tb6_full = slide6.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))
    tf6_full = tb6_full.text_frame
    tf6_full.word_wrap = True

    sections6 = [
        ("• Details / Links of the reference and research work:", [
            "JCERT Primary Textbooks ('भाषा अंजलि', 'गणित ज्ञान', 'हमारा परिवेश'): Jharkhand Council of Educational Research and Training.",
            "PALASH MTB-MLE Programme Reports (2022-2025): Dept. of School Education & Literacy, Govt. of Jharkhand.",
            "National Education Policy (NEP 2020) §4.11-4.14: Mandating mother-tongue / home language based primary instruction & continuous FLN assessment.",
            "NIPUN Bharat & Vidya Pravesh Guidelines: Ministry of Education, Govt. of India (Foundational Literacy and Numeracy Assessment Framework).",
            "Central Institute of Indian Languages (CIIL Mysore): Linguistic research on Austroasiatic (Munda) language family syntax and phonology.",
            "Guru Gomke Pt. Raghunath Murmu & Lako Bodra: Script mapping standards for Ol Chiki (ᱚᱞ ᱪᱤᱠᱤ) and Warang Chiti.",
            "W3C Web Speech & Web Audio API Standards: Low-latency client-side speech recognition & acoustic synthesis on constrained edge devices.",
            "Working Prototype Repository: Fully functional BHASHASETU PWA + FastAPI Server with 100% offline capabilities & NIPUN Assessment Exam Suite."
        ])
    ]

    p = tf6_full.paragraphs[0]
    p.space_after = Pt(6)
    r = p.add_run()
    r.text = sections6[0][0]
    r.font.name = 'Arial'
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = HEADING_BLACK

    for item in sections6[0][1]:
        pi = tf6_full.add_paragraph()
        pi.space_after = Pt(6)
        r_dot = pi.add_run()
        r_dot.text = "   • "
        r_dot.font.bold = True
        r_dot.font.size = Pt(12)
        r_dot.font.color.rgb = PALASH_ORANGE

        r_txt = pi.add_run()
        r_txt.text = item
        r_txt.font.size = Pt(12)
        r_txt.font.color.rgb = TEXT_BLACK

    # Save presentation safely
    try:
        prs.save(output_path)
        print(f"Presentation saved successfully to {output_path}")
    except PermissionError:
        alt_path = "SIH2026_Idea_Submission_PS26042_Latest.pptx"
        prs.save(alt_path)
        print(f"File was locked by viewer, saved successfully to {alt_path}")

if __name__ == "__main__":
    build_presentation()
